"""업로드한 문서를 텍스트로 읽고, 사용자가 확인한 부분을 메모로 저장한다.

원본 파일은 디스크에 저장하지 않는다. 페이지·시트·구역별 텍스트만 미리보기로
돌려주며, 저장은 별도 요청으로 수행한다. HWP 5.x는 한컴 공개 형식에 따른다.
"""

from __future__ import annotations

import hashlib
import io
import posixpath
import re
import struct
import zipfile
import zlib
from datetime import date, datetime
from pathlib import PurePosixPath

from . import db, store, textutil
from .importers import clean

MAX_BYTES = 20_000_000
MAX_TEXT = 1_000_000
MAX_PARTS = 500
MAX_UNPACKED = 50_000_000
MAX_CELLS = 200_000
FORMATS = {
    ".pdf": "PDF", ".xlsx": "엑셀", ".xlsm": "엑셀", ".xls": "엑셀",
    ".hwp": "한글", ".hwpx": "한글", ".txt": "텍스트", ".md": "텍스트",
}


class DocumentError(ValueError):
    """화면에 그대로 표시할 수 있는 문서 읽기 오류."""


def filename_of(value: str) -> str:
    if not isinstance(value, str):
        raise DocumentError("파일 이름을 확인하세요.")
    name = clean(value.replace("\\", "/").rsplit("/", 1)[-1])
    if not name or len(name) > 200 or "\n" in name or "\t" in name:
        raise DocumentError("파일 이름은 200자 이하로 지정하세요.")
    if PurePosixPath(name).suffix.lower() not in FORMATS:
        raise DocumentError("PDF, 엑셀(.xlsx·.xlsm·.xls), 한글(.hwp·.hwpx), TXT·MD 파일을 선택하세요.")
    return name


def extract(filename: str, data: bytes) -> dict:
    name = filename_of(filename)
    if not data:
        raise DocumentError("파일이 비어 있습니다.")
    if len(data) > MAX_BYTES:
        raise DocumentError("20MB 이하의 파일을 선택하세요.")
    suffix = PurePosixPath(name).suffix.lower()
    readers = {".pdf": _pdf, ".xlsx": _xlsx, ".xlsm": _xlsx, ".xls": _xls,
               ".hwp": _hwp, ".hwpx": _hwpx, ".txt": _text, ".md": _text}
    warnings: list[str] = []
    try:
        parts = readers[suffix](data, warnings)
    except DocumentError:
        raise
    except ImportError:
        raise DocumentError("이 설치 버전에는 해당 문서를 읽는 구성요소가 없습니다. 문서 가져오기 설치 안내를 확인하세요.") from None
    except Exception as exc:
        raise DocumentError("파일을 읽지 못했습니다. 파일이 손상되었거나 지원하지 않는 형식일 수 있습니다.") from exc
    parts = [{**p, "text": clean(p["text"])} for p in parts if clean(p["text"])]
    if not parts:
        raise DocumentError("읽을 수 있는 텍스트가 없습니다. 스캔·이미지 문서는 OCR 처리 후 가져오세요.")
    total = sum(len(p["text"]) for p in parts)
    if total > MAX_TEXT or len(parts) > MAX_PARTS:
        raise DocumentError("문서가 너무 큽니다. 100만 자·500개 부분 이하로 나눠 가져오세요.")
    return {"filename": name, "format": FORMATS[suffix], "parts": parts,
            "warnings": list(dict.fromkeys(warnings)), "chars": total}


def _part(key: str, label: str, text: str) -> dict:
    return {"key": key, "label": label, "text": text}


def _archive(data: bytes) -> zipfile.ZipFile:
    archive = zipfile.ZipFile(io.BytesIO(data))
    infos = archive.infolist()
    if len(infos) > 10000 or sum(i.file_size for i in infos) > MAX_UNPACKED:
        archive.close()
        raise DocumentError("압축을 푼 문서가 너무 큽니다. 파일을 나눠 가져오세요.")
    if any(i.flag_bits & 1 for i in infos):
        archive.close()
        raise DocumentError("암호가 걸린 문서는 읽을 수 없습니다. 암호를 해제한 사본을 선택하세요.")
    return archive


def _pdf(data: bytes, warnings: list[str]) -> list[dict]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted and not reader.decrypt(""):
        raise DocumentError("암호가 걸린 PDF입니다. 암호를 해제한 사본을 선택하세요.")
    if len(reader.pages) > MAX_PARTS:
        raise DocumentError("PDF는 500쪽 이하로 나눠 가져오세요.")
    parts, missing, total = [], [], 0
    for i, page in enumerate(reader.pages, 1):
        content = page.get_contents()
        if content is not None and len(content.get_data()) > MAX_UNPACKED:
            raise DocumentError(f"PDF {i}쪽이 너무 큽니다. 해당 페이지를 나눠 가져오세요.")
        text = clean(page.extract_text() or "")
        total += len(text)
        if total > MAX_TEXT:
            raise DocumentError("PDF 텍스트가 100만 자를 넘습니다. 문서를 나눠 가져오세요.")
        if text:
            parts.append(_part(f"page:{i}", f"{i}쪽", text))
        else:
            missing.append(str(i))
    if missing:
        warnings.append(f"텍스트를 읽지 못한 페이지: {', '.join(missing[:30])}쪽"
                        f"{' 외 ' + str(len(missing) - 30) + '쪽' if len(missing) > 30 else ''}. 스캔·이미지 페이지는 OCR이 필요합니다.")
    warnings.append("PDF의 표·단 배치는 원본과 다를 수 있습니다. 미리보기에서 내용을 확인하세요.")
    return parts


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat()
    return str(value).replace("\t", " ").replace("\n", " / ")


def _xlsx(data: bytes, warnings: list[str]) -> list[dict]:
    from openpyxl import load_workbook

    warnings.append("엑셀의 서식·차트·이미지는 제외하고 셀의 저장값을 읽습니다.")
    with _archive(data):
        pass
    values = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
    formulas = None
    try:
        formulas = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        parts, cells, total = [], 0, 0
        if len(values.worksheets) > MAX_PARTS:
            raise DocumentError("엑셀은 500개 시트 이하로 나눠 가져오세요.")
        for sheet in values.worksheets:
            if sheet.sheet_state != "visible":
                warnings.append("숨겨진 시트는 가져오지 않았습니다.")
                continue
            # 잘못 저장된 dimension 때문에 본문을 놓치지 않도록 실제 XML 행을 읽는다.
            sheet.reset_dimensions()
            formula_sheet = formulas[sheet.title]
            formula_sheet.reset_dimensions()
            lines = []
            for i, (row, formula_row) in enumerate(zip(sheet.iter_rows(), formula_sheet.iter_rows()), 1):
                cells += len(row)
                if cells > MAX_CELLS:
                    raise DocumentError("엑셀의 셀이 너무 많습니다. 20만 셀 이하로 나눠 가져오세요.")
                out = []
                for cell, formula in zip(row, formula_row):
                    value = cell.value
                    if value is None and formula.data_type == "f":
                        value = formula.value
                        warnings.append("저장된 계산 결과가 없는 셀은 수식으로 표시했습니다. 앱에서 수식을 계산하지는 않습니다.")
                    out.append(_cell(value))
                while out and not out[-1]:
                    out.pop()
                if not any(out):
                    continue
                line = f"{i}행: " + "\t".join(out)
                total += len(line)
                if total > MAX_TEXT:
                    raise DocumentError("엑셀 텍스트가 100만 자를 넘습니다. 시트를 나눠 가져오세요.")
                lines.append(line)
            if lines:
                parts.append(_part(f"sheet:{sheet.title}", f"시트: {sheet.title}", "\n".join(lines)))
        return parts
    finally:
        values.close()
        if formulas is not None:
            formulas.close()


def _xls(data: bytes, warnings: list[str]) -> list[dict]:
    import xlrd

    warnings.append("엑셀의 서식·차트·이미지는 제외하고 셀의 저장값을 읽습니다.")
    book = xlrd.open_workbook(file_contents=data, on_demand=True)
    try:
        parts, cells, total = [], 0, 0
        if book.nsheets > MAX_PARTS:
            raise DocumentError("엑셀은 500개 시트 이하로 나눠 가져오세요.")
        for sheet in book.sheets():
            if sheet.visibility:
                warnings.append("숨겨진 시트는 가져오지 않았습니다.")
                continue
            lines = []
            for r in range(sheet.nrows):
                cells += sheet.ncols
                if cells > MAX_CELLS:
                    raise DocumentError("엑셀의 셀이 너무 많습니다. 20만 셀 이하로 나눠 가져오세요.")
                values = []
                for cell in sheet.row(r):
                    value = cell.value
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        value = xlrd.xldate_as_datetime(value, book.datemode)
                    elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        value = bool(value)
                    values.append(_cell(value))
                while values and not values[-1]:
                    values.pop()
                if any(values):
                    line = f"{r + 1}행: " + "\t".join(values)
                    total += len(line)
                    if total > MAX_TEXT:
                        raise DocumentError("엑셀 텍스트가 100만 자를 넘습니다. 시트를 나눠 가져오세요.")
                    lines.append(line)
            if lines:
                parts.append(_part(f"sheet:{sheet.name}", f"시트: {sheet.name}", "\n".join(lines)))
        return parts
    finally:
        book.release_resources()


def _xml(data: bytes):
    from defusedxml.ElementTree import fromstring

    return fromstring(data)


def _local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _hwpx(data: bytes, warnings: list[str]) -> list[dict]:
    with _archive(data) as archive:
        names = archive.namelist()
        if "META-INF/manifest.xml" in names:
            manifest = _xml(archive.read("META-INF/manifest.xml"))
            if any("encryption" in _local(e.tag).lower() for e in manifest.iter()):
                raise DocumentError("암호가 걸린 HWPX입니다. 암호를 해제한 사본을 선택하세요.")
        sections = sorted((n for n in names if re.fullmatch(r"Contents/section\d+\.xml", n)),
                          key=lambda n: int(re.search(r"section(\d+)", n)[1]))
        if "Contents/content.hpf" in names:
            package = _xml(archive.read("Contents/content.hpf"))
            mapping = {e.get("id"): posixpath.normpath(posixpath.join("Contents", e.get("href", "")))
                       for e in package.iter() if _local(e.tag) == "item"}
            order = [mapping.get(e.get("idref")) for e in package.iter() if _local(e.tag) == "itemref"]
            ordered = [n for n in order if n in sections]
            sections = list(dict.fromkeys(ordered + sections))
        if not sections:
            raise DocumentError("HWPX 본문을 찾지 못했습니다. 파일 형식을 확인하세요.")
        if len(sections) > MAX_PARTS:
            raise DocumentError("한글 문서는 500개 구역 이하로 나눠 가져오세요.")
        parts, total = [], 0
        for i, name in enumerate(sections, 1):
            output = []

            def text_node(node):
                output.append(node.text or "")
                for child in node:
                    if _local(child.tag) in {"lineBreak", "br"}:
                        output.append("\n")
                    elif _local(child.tag) == "tab":
                        output.append("\t")
                    else:
                        text_node(child)
                    output.append(child.tail or "")

            def walk(node):
                if _local(node.tag) == "t":
                    text_node(node)
                    return
                for child in node:
                    walk(child)
                if _local(node.tag) == "p":
                    output.append("\n")

            walk(_xml(archive.read(name)))
            text = clean("".join(output))
            total += len(text)
            if total > MAX_TEXT:
                raise DocumentError("한글 문서 텍스트가 100만 자를 넘습니다. 문서를 나눠 가져오세요.")
            if text:
                parts.append(_part(name, f"{i}번째 구역", text))
        warnings.append("한글 문서의 글자와 표 안의 텍스트를 읽습니다. 그림·필기·원본 배치는 저장하지 않습니다.")
        return parts


def _hwp_paragraph(data: bytes) -> str:
    if len(data) % 2:
        raise DocumentError("HWP 문단이 손상되어 읽지 못했습니다.")
    output, pos = [], 0
    while pos < len(data):
        code = struct.unpack_from("<H", data, pos)[0]
        if 1 <= code <= 23 and code not in {10, 13}:
            if pos + 16 > len(data):
                raise DocumentError("HWP 제어 정보가 손상되어 읽지 못했습니다.")
            if code == 9:
                output.append(b"\t\x00")
            pos += 16
            continue
        if code in {10, 13}:
            output.append(b"\n\x00")
        elif code == 24:
            output.append(b"-\x00")
        elif code in {30, 31}:
            output.append(b" \x00")
        elif code >= 32:
            output.append(data[pos:pos + 2])
        pos += 2
    return b"".join(output).decode("utf-16le")


def _hwp_section(data: bytes) -> str:
    pos, paragraphs, total = 0, [], 0
    while pos < len(data):
        if pos + 4 > len(data):
            raise DocumentError("HWP 본문 레코드가 손상되어 읽지 못했습니다.")
        header = struct.unpack_from("<I", data, pos)[0]
        pos += 4
        size = header >> 20
        if size == 0xFFF:
            if pos + 4 > len(data):
                raise DocumentError("HWP 본문 길이 정보가 손상되었습니다.")
            size = struct.unpack_from("<I", data, pos)[0]
            pos += 4
        if pos + size > len(data):
            raise DocumentError("HWP 본문 길이 정보가 손상되었습니다.")
        if header & 0x3FF == 0x43:  # HWPTAG_PARA_TEXT
            text = _hwp_paragraph(data[pos:pos + size])
            total += len(text)
            if total > MAX_TEXT:
                raise DocumentError("한글 문서 텍스트가 100만 자를 넘습니다.")
            paragraphs.append(text)
        pos += size
    return clean("\n".join(paragraphs))


def _hwp(data: bytes, warnings: list[str]) -> list[dict]:
    import olefile

    if not olefile.isOleFile(io.BytesIO(data)):
        raise DocumentError("HWP 5.x 파일이 아닙니다. 한글에서 HWPX로 다시 저장해 가져오세요.")
    with olefile.OleFileIO(io.BytesIO(data)) as document:
        header = document.openstream("FileHeader").read()
        if len(header) < 40 or not header.startswith(b"HWP Document File") or header[35] != 5:
            raise DocumentError("지원하지 않는 HWP 형식입니다. HWPX로 다시 저장해 가져오세요.")
        flags = struct.unpack_from("<I", header, 36)[0]
        if flags & ((1 << 1) | (1 << 2) | (1 << 4) | (1 << 8) | (1 << 10)):
            raise DocumentError("암호·보안·배포용 HWP는 지원하지 않습니다. 읽을 수 있는 HWPX나 PDF 사본을 선택하세요.")
        sections = sorted((path for path in document.listdir()
                           if len(path) == 2 and path[0] == "BodyText" and re.fullmatch(r"Section\d+", path[1])),
                          key=lambda path: int(path[1][7:]))
        if not sections or len(sections) > MAX_PARTS:
            raise DocumentError("HWP 본문을 찾지 못했거나 구역이 너무 많습니다. HWPX로 저장해 가져오세요.")
        parts, total = [], 0
        for i, path in enumerate(sections, 1):
            if document.get_size(path) > MAX_UNPACKED:
                raise DocumentError("HWP 구역이 너무 큽니다. 문서를 나눠 가져오세요.")
            raw = document.openstream(path).read()
            if flags & 1:
                inflater = zlib.decompressobj(-15)
                raw = inflater.decompress(raw, MAX_UNPACKED + 1)
                if len(raw) > MAX_UNPACKED or inflater.unconsumed_tail or not inflater.eof:
                    raise DocumentError("HWP 압축 본문이 너무 크거나 손상되었습니다.")
            text = _hwp_section(raw)
            total += len(text)
            if total > MAX_TEXT:
                raise DocumentError("한글 문서 텍스트가 100만 자를 넘습니다. 문서를 나눠 가져오세요.")
            if text:
                parts.append(_part("/".join(path), f"{i}번째 구역", text))
        warnings.append("한글 문서의 글자와 표 안의 텍스트를 읽습니다. 그림·필기·원본 배치는 저장하지 않습니다.")
        return parts


def _text(data: bytes, warnings: list[str]) -> list[dict]:
    encodings = ["utf-16"] if data.startswith((b"\xff\xfe", b"\xfe\xff")) else ["utf-8-sig", "cp949", "utf-16"]
    for encoding in encodings:
        try:
            return [_part("text", "본문", data.decode(encoding))]
        except UnicodeError:
            continue
    raise DocumentError("텍스트 인코딩을 읽지 못했습니다. UTF-8 텍스트로 저장해 가져오세요.")


def save(filename: str, parts: list[dict]) -> dict:
    """편집한 미리보기를 한 번에 저장한다. 동일 파일·위치·본문은 중복되지 않는다."""
    name = filename_of(filename)
    if not isinstance(parts, list) or not 1 <= len(parts) <= MAX_PARTS:
        raise DocumentError("저장할 부분을 1개 이상, 500개 이하로 선택하세요.")
    validated, keys, total = [], set(), 0
    for part in parts:
        if not isinstance(part, dict) or any(not isinstance(part.get(k), str) for k in ("key", "label", "text")):
            raise DocumentError("저장할 텍스트와 원본 위치를 확인하세요.")
        key, label, text = part["key"], clean(part["label"]), clean(part["text"])
        if not key or len(key) > 300 or key in keys or not label or len(label) > 200 or not text:
            raise DocumentError("선택한 부분의 내용이나 원본 위치가 비어 있거나 중복됩니다.")
        keys.add(key)
        total += len(text)
        if total > MAX_TEXT:
            raise DocumentError("저장할 텍스트가 100만 자를 넘습니다. 나눠 저장하세요.")
        uid = hashlib.sha256(f"{name}\0{key}\0{text}".encode("utf-8")).hexdigest()
        kind = FORMATS[PurePosixPath(name).suffix.lower()]
        body = f"{name} · {label}\n원본 파일: {name}\n원본 위치: {label}\n\n{text}\n\n#문서 #{kind}"
        validated.append((uid, body))

    conn, created, memos = db.connect(), [], []
    now = store.now_iso()
    with conn:
        conn.execute("BEGIN IMMEDIATE")
        for uid, body in validated:
            existing = conn.execute("SELECT id, title FROM memos WHERE source='document' AND external_id=?", (uid,)).fetchone()
            if existing:
                memos.append({"id": existing["id"], "title": existing["title"], "duplicate": True})
                continue
            title = textutil.title_from(body)
            cursor = conn.execute(
                "INSERT INTO memos(body,title,source,created_at,updated_at,external_id) VALUES(?,?,'document',?,?,?)",
                (body, title, now, now, uid),
            )
            memo_id = cursor.lastrowid
            conn.execute("INSERT INTO memo_fts(rowid,grams,raw) VALUES(?,?,?)",
                         (memo_id, textutil.ngram_index_text(body), textutil.normalize(body).lower()))
            created.append(memo_id)
            memos.append({"id": memo_id, "title": title, "duplicate": False})
    for memo_id in created:
        store.enqueue(memo_id)
    return {"imported": len(created), "skipped_duplicate": len(memos) - len(created), "memos": memos}
