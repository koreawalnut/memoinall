"""실제 파서를 통과하는 작은 문서 파일들. 사용자 파일·네트워크 없이 생성한다."""

import io
import struct
import zipfile
import zlib


def xlsx_bytes():
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = "인수인계"
    sheet.append(["업무", "담당자", "금액"])
    sheet.append(["결제 점검", "김민수", 1200])
    sheet.append(["수식", "", "=SUM(C2:C2)"])
    book.create_sheet("참고").append(["배포 전 로그 확인"])
    hidden = book.create_sheet("숨김")
    hidden.sheet_state = "hidden"
    hidden.append(["숨긴 자료"])
    out = io.BytesIO()
    book.save(out)
    book.close()
    return out.getvalue()


def xls_bytes():
    import xlwt

    book = xlwt.Workbook()
    sheet = book.add_sheet("업무")
    sheet.write(0, 0, "구형 엑셀 인수인계")
    sheet.write(0, 1, 1200)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def pdf_bytes(*, encrypted=False, blank=False):
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    if not blank:
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                                 NameObject("/Subtype"): NameObject("/Type1"),
                                 NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
        content = DecodedStreamObject()
        content.set_data(b"BT /F1 12 Tf 50 700 Td (Project handoff: payment checks) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(content)
        writer.add_blank_page(width=612, height=792)
    if encrypted:
        writer.encrypt("test-password")
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def hwpx_bytes(*, malicious=False, encrypted=False):
    section = ('<hs:sec xmlns:hs="http://www.hancom.co.kr/hwpml/2011/section" '
               'xmlns:hp="http://www.hancom.co.kr/hwpml/2011/paragraph">'
               '<hp:p><hp:run><hp:t>한글 업무 인수인계</hp:t><hp:t> 자료</hp:t></hp:run></hp:p>'
               '<hp:p><hp:run><hp:t>다음 확인</hp:t></hp:run></hp:p></hs:sec>')
    if malicious:
        section = '<!DOCTYPE sec [<!ENTITY x "untrusted">]><sec>&x;</sec>'
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("mimetype", "application/hwp+zip")
        archive.writestr("Contents/section0.xml", section)
        archive.writestr("Contents/section1.xml", '<sec><p><t>먼저 읽을 구역</t></p></sec>')
        archive.writestr("Contents/content.hpf", '<package><manifest><item id="a" href="section0.xml"/>'
                         '<item id="b" href="section1.xml"/></manifest><spine><itemref idref="b"/>'
                         '<itemref idref="a"/></spine></package>')
        if encrypted:
            archive.writestr("META-INF/manifest.xml", '<manifest><encryption-data/></manifest>')
    return out.getvalue()


def hwp_bytes(*, flags=1, broken=False):
    """최소 CFB: FileHeader, BodyText/Section0. 모든 스트림은 regular FAT에 둔다."""
    free, end, fat = 0xFFFFFFFF, 0xFFFFFFFE, 0xFFFFFFFD
    file_header = bytearray(4096)
    file_header[:17] = b"HWP Document File"
    struct.pack_into("<II", file_header, 32, 0x05000300, flags)
    # 확장 제어(표), 탭, UTF-16 surrogate pair를 포함한다.
    paragraph = "한글 본문 ".encode("utf-16le") + struct.pack("<H", 11) + b"TABLE-CONTROL!"
    paragraph += "표 내용 😀".encode("utf-16le") + struct.pack("<H", 9) + b"\x00" * 14 + "다음\r".encode("utf-16le")
    record = struct.pack("<I", 0x43 | (1 << 10) | (len(paragraph) << 20)) + paragraph
    if broken:
        record = struct.pack("<I", 0x43 | (2000 << 20)) + b"x"
    if flags & 1:
        compressor = zlib.compressobj(wbits=-15)
        record = compressor.compress(record) + compressor.flush()
    body = record.ljust(4096, b"\x00")

    header = bytearray(512)
    header[:8] = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
    struct.pack_into("<HHHHH", header, 24, 0x3E, 3, 0xFFFE, 9, 6)
    struct.pack_into("<IIIIIIIII", header, 40, 0, 1, 1, 0, 4096, end, 0, end, 0)
    struct.pack_into("<109I", header, 76, 0, *([free] * 108))
    allocation = [free] * 128
    allocation[0], allocation[1] = fat, end
    for start in (2, 10):
        for sector in range(start, start + 7):
            allocation[sector] = sector + 1
        allocation[start + 7] = end

    def entry(name, kind, *, child=free, right=free, start=end, size=0):
        item = bytearray(128)
        encoded = (name + "\0").encode("utf-16le")
        item[:len(encoded)] = encoded
        struct.pack_into("<HBBIII", item, 64, len(encoded), kind, 1, free, right, child)
        struct.pack_into("<IQ", item, 116, start, size)
        return bytes(item)

    directory = (entry("Root Entry", 5, child=1) + entry("FileHeader", 2, right=2, start=2, size=4096)
                 + entry("BodyText", 1, child=3) + entry("Section0", 2, start=10, size=4096))
    return bytes(header) + struct.pack("<128I", *allocation) + directory + bytes(file_header) + body
